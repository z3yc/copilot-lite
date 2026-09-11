"""解析器与分块测试。"""


from app.rag.chunking import chunk_document
from app.rag.parsers import get_parser

_MD_SAMPLE = """# 项目简介

这是一个 AI 助理项目。

## 技术栈

- FastAPI
- DeepSeek

## 架构

分为前端与后端。
"""


def test_markdown_parser_sections() -> None:
    """Markdown 按标题层级切分。"""
    parsed = get_parser("md").parse(_MD_SAMPLE.encode())
    assert parsed.title == "项目简介"
    headings = [s.heading for s in parsed.sections]
    assert headings == ["项目简介", "技术栈", "架构"]
    assert "FastAPI" in parsed.sections[1].content


def test_markdown_parser_plain_text() -> None:
    """无标题的纯文本整体成段。"""
    parsed = get_parser("md").parse("这是没有标题的笔记内容，仅此而已。".encode())
    assert len(parsed.sections) == 1
    assert parsed.sections[0].heading is None


def test_chunking_heading_path() -> None:
    """分块携带完整标题路径（溯源用）。"""
    content = """# 第一章

第一章内容。

## 第一节

第一节内容，用于检索测试。
"""
    parsed = get_parser("md").parse(content.encode())
    chunks = chunk_document(parsed)
    assert chunks
    # 含"第一节"的块应携带标题路径 ["第一章", "第一节"]
    target = next(c for c in chunks if "第一节内容" in c.content)
    assert target.meta["headings"] == ["第一章", "第一节"]


def test_chunking_long_text_split() -> None:
    """超长内容按段落切分为多块，且带重叠。"""
    long_body = "\n\n".join(f"这是第{i}段的内容，包含一些检索关键词。" for i in range(30))
    parsed = get_parser("md").parse(f"# 标题\n\n{long_body}".encode())
    chunks = chunk_document(parsed, max_size=100, overlap=20)
    assert len(chunks) > 1
    # 所有块都有 chunk_index 且连续
    indexes = [c.meta["chunk_index"] for c in chunks]
    assert indexes == list(range(len(chunks)))


def test_chunking_overlap_at_aggregation_boundary() -> None:
    """聚合溢出边界也有 overlap（回归：此前仅超长单段落有重叠）。"""
    p1 = "甲" * 60
    p2 = "乙" * 60
    parsed = get_parser("md").parse(f"# 标题\n\n{p1}\n\n{p2}".encode())
    chunks = chunk_document(parsed, max_size=100, overlap=20)
    assert len(chunks) == 2
    assert chunks[1].content.startswith("甲" * 20), "第二块应以第一块结尾 20 字符开头"


def test_parser_registry() -> None:
    """五种解析器全部注册。"""
    types = sorted(get_parser(t) is not None for t in ["md", "pdf", "docx", "code", "web"])
    assert all(types)


def test_code_parser() -> None:
    """代码解析器：保留文件路径信息。"""
    parsed = get_parser("code").parse(
        b"def hello():\n    return 'hi'\n",
        meta={"file_path": "src/main.py"},
    )
    assert parsed.title == "src/main.py"
    assert parsed.sections[0].heading == "src/main.py"
    assert "hello" in parsed.sections[0].content


def test_web_parser() -> None:
    """网页解析器：提取标题结构与正文，剔除噪声。"""
    html = """<html><head><title>测试页面</title></head>
    <body>
      <script>var x = 1;</script>
      <nav>导航</nav>
      <h1>主标题</h1>
      <p>这是正文第一段。</p>
      <h2>小标题</h2>
      <p>第二段内容。</p>
    </body></html>"""
    parsed = get_parser("web").parse(html.encode())
    assert parsed.title == "测试页面"
    headings = [s.heading for s in parsed.sections]
    assert headings == ["主标题", "小标题"]
    assert "第一段" in parsed.sections[0].content
    assert "导航" not in parsed.sections[0].content  # nav 已被剔除


def test_docx_parser() -> None:
    """DOCX 解析器：按 Heading 样式切分。"""
    from io import BytesIO

    from docx import Document as DocxDocument

    doc = DocxDocument()
    doc.add_heading("标题一", level=1)
    doc.add_paragraph("标题一的内容。")
    doc.add_heading("标题二", level=2)
    doc.add_paragraph("标题二的内容。")
    buf = BytesIO()
    doc.save(buf)

    parsed = get_parser("docx").parse(buf.getvalue())
    headings = [s.heading for s in parsed.sections]
    assert headings == ["标题一", "标题二"]
    assert "标题一的内容" in parsed.sections[0].content


def test_pdf_parser_blank() -> None:
    """PDF 解析器：无文本页不产生章节（空白 PDF）。"""
    from io import BytesIO

    from pypdf import PdfWriter

    writer = PdfWriter()
    writer.add_blank_page(width=200, height=200)
    buf = BytesIO()
    writer.write(buf)

    parsed = get_parser("pdf").parse(buf.getvalue())
    assert parsed.source_type == "pdf"
    assert parsed.sections == []


def test_decode_text_handles_utf16_and_strips_nul():
    """UTF-16/含 NUL 文本：解码后不得残留 NUL（PostgreSQL 不接受）。"""
    from app.rag.parsers.base import decode_text

    raw = "hello ==1\n".encode("utf-16")  # 带 BOM + 字节间 \x00
    text = decode_text(raw)
    assert "\x00" not in text
    assert "hello ==1" in text
    assert decode_text(b"a\x00b") == "ab"
    assert decode_text(b"") == ""


def test_markdown_parser_handles_utf16():
    from app.rag.parsers.markdown_parser import MarkdownParser

    parsed = MarkdownParser().parse("# 标题\n\n正文".encode("utf-16"))
    joined = "".join(s.content for s in parsed.sections)
    assert "\x00" not in joined
    assert "正文" in joined
