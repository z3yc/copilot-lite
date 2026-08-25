"""DOCX 解析器：按 Word 内置 Heading 样式切分章节。"""

from io import BytesIO

from app.rag.parsers.base import DocumentParser, ParsedDocument, Section, register_parser


class DocxParser(DocumentParser):
    source_type = "docx"

    def parse(self, content: bytes, meta: dict | None = None) -> ParsedDocument:
        from docx import Document as DocxDocument

        meta = meta or {}
        doc = DocxDocument(BytesIO(content))

        sections: list[Section] = []
        current_heading: str | None = None
        current_level = 0
        buffer: list[str] = []

        def flush() -> None:
            body = "\n".join(buffer).strip()
            if body or current_heading is not None:
                sections.append(
                    Section(heading=current_heading, level=current_level, content=body)
                )
            buffer.clear()

        for para in doc.paragraphs:
            style = para.style.name if para.style else ""
            text = para.text.strip()
            if not text:
                continue
            if style.startswith("Heading"):
                flush()
                current_heading = text
                current_level = int(style.replace("Heading ", "").split()[0]) if style != "Heading" else 1
            elif style == "Title":
                flush()
                current_heading = text
                current_level = 1
            else:
                buffer.append(text)

        flush()

        title = meta.get("title") or (
            doc.core_properties.title or "未命名DOCX"
        )
        return ParsedDocument(
            title=title, source_type=self.source_type, sections=sections, meta=meta
        )


register_parser(DocxParser())
